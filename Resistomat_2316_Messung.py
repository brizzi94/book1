"""
RESISTOMAT 2316 - manuelle Messung mit UI
------------------------------------------
Ablauf:
1. Beim Start: Anzahl Messungen eingeben, dann Speicherort fuer die
   Excel-Datei per Dialog waehlen.
2. Das Fenster zeigt fortlaufend den aktuellen Messwert vom Geraet an
   (wird im Hintergrund alle 0.3s abgefragt).
3. Du misst manuell am Geraet. Sobald der angezeigte Wert passt,
   druecke ENTER (oder klicke "Bestaetigen") - der aktuell angezeigte
   Wert wird mit Zeitstempel in die Excel-Datei geschrieben.
4. Nach N Messungen ("Anzahl") ist Schluss, das Fenster meldet "Fertig".

Vorher pruefen:
- pip install pyserial openpyxl
- Am Geraet, Menue 150 (RS232): BLOCKCHECK muss auf OFF stehen
- COM-Port wird beim Start im GUI ausgewaehlt, BAUDRATE unten anpassen
"""

import json
import re
import time
import datetime
import threading
import queue
import tkinter as tk
from tkinter import simpledialog, filedialog, messagebox, ttk

import serial
import serial.tools.list_ports
from pathlib import Path
from openpyxl import Workbook, load_workbook

BAUDRATE = 9600        # <-- muss mit Geraete-Einstellung uebereinstimmen
POLL_INTERVAL = 0.3

# Merkt sich den zuletzt gewaehlten Speicherort/COM-Port, damit die
# Dialoge beim naechsten Start dieselben Werte vorschlagen.
CONFIG_PATH = Path.home() / ".resistomat_2316_config.json"


def load_config():
    try:
        return json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
    except (FileNotFoundError, json.JSONDecodeError, OSError):
        return {}


def save_config_value(key: str, value: str):
    data = load_config()
    data[key] = value
    try:
        CONFIG_PATH.write_text(json.dumps(data), encoding="utf-8")
    except OSError:
        pass


def load_last_excel_path():
    return load_config().get("last_excel_path")


def save_last_excel_path(path: str):
    save_config_value("last_excel_path", path)


def load_last_com_port():
    return load_config().get("last_com_port")


def save_last_com_port(port: str):
    save_config_value("last_com_port", port)

STX = b"\x02"
ETX = b"\x03"
ENQ = b"\x05"
ACK = b"\x06"
NAK = b"\x15"
LF = b"\x0A"
EOT = b"\x04"
ADDR = b"0000"


# ---------------------------------------------------------
# RESISTOMAT-Kommunikation (wie im bisherigen Skript)
# ---------------------------------------------------------
def send(ser, command: str) -> bool:
    frame = EOT + ADDR + b"sr" + STX + command.encode("ascii") + LF + ETX
    ser.write(frame)
    return ser.read(1) == ACK


def poll(ser):
    ser.write(EOT + ADDR + b"po" + ENQ)
    raw = ser.read_until(ETX)
    if STX not in raw:
        return None
    ser.write(ACK)
    ser.read(1)  # trailing EOT
    return raw.split(STX, 1)[1].split(ETX, 1)[0].decode("ascii", errors="replace").strip()


def fetch_value(ser):
    if not send(ser, "fetc?"):
        return None
    return poll(ser)


def ask_com_port(root, default_port=None):
    available = [p.device for p in serial.tools.list_ports.comports()]

    dialog = tk.Toplevel(root)
    dialog.title("COM-Port waehlen")
    dialog.geometry("300x150")
    dialog.resizable(False, False)
    dialog.transient(root)
    dialog.grab_set()

    tk.Label(dialog, text="Serieller Port (COM):", font=("Segoe UI", 11)).pack(pady=(15, 5))

    initial = default_port or (available[0] if available else "COM1")
    port_var = tk.StringVar(value=initial)
    combo = ttk.Combobox(dialog, textvariable=port_var, values=available, width=20)
    combo.pack(pady=5)
    combo.focus_set()

    result = {"port": None}

    def on_ok():
        result["port"] = port_var.get().strip()
        dialog.destroy()

    def on_cancel():
        dialog.destroy()

    button_frame = tk.Frame(dialog)
    button_frame.pack(pady=15)
    tk.Button(button_frame, text="Verbinden", command=on_ok, width=10).pack(side="left", padx=5)
    tk.Button(button_frame, text="Abbrechen", command=on_cancel, width=10).pack(side="left", padx=5)

    dialog.bind("<Return>", lambda event: on_ok())
    dialog.protocol("WM_DELETE_WINDOW", on_cancel)
    dialog.wait_window()
    return result["port"] or None


def split_value_unit(raw_value: str):
    match = re.match(r"([-+]?\d*\.?\d+)\s*([A-Za-z]*)", raw_value)
    if not match:
        return None
    number = float(match.group(1))
    unit = match.group(2)
    return number, unit


def ensure_excel_file(path: str):
    file = Path(path)
    if file.exists():
        wb = load_workbook(path)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Messwerte"
        ws.append(["Datum/Zeit", "Messwert", "Einheit"])
    return wb, ws


def append_measurement(wb, ws, path: str, number: float, unit: str):
    """Haengt eine Messung an und gibt die Zeilennummer zurueck, oder
    None wenn das Speichern fehlschlug (z.B. Datei in Excel geoeffnet)."""
    ws.append([datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"), number, unit])
    try:
        wb.save(path)
        return ws.max_row
    except PermissionError:
        ws.delete_rows(ws.max_row)
        return None


# ---------------------------------------------------------
# Hintergrund-Thread: fragt laufend den aktuellen Wert ab
# ---------------------------------------------------------
def polling_worker(ser, value_queue, stop_event):
    while not stop_event.is_set():
        raw_value = fetch_value(ser)
        if raw_value is not None:
            value_queue.put(raw_value)
        time.sleep(POLL_INTERVAL)


# ---------------------------------------------------------
# UI
# ---------------------------------------------------------
class ResistomatUI:
    def __init__(self, root, ser, excel_path, target_count):
        self.root = root
        self.ser = ser
        self.excel_path = excel_path
        self.target_count = target_count
        self.count = 0
        self.current_raw_value = None
        self.history_rows = []   # Excel-Zeilennummern, Index 0 = neuester Eintrag
        self.history_raw = []    # Rohwerte (Text), gleiche Reihenfolge wie history_rows

        self.wb, self.ws = ensure_excel_file(excel_path)

        self.value_queue = queue.Queue()
        self.stop_event = threading.Event()
        self.worker_thread = threading.Thread(
            target=polling_worker,
            args=(ser, self.value_queue, self.stop_event),
            daemon=True,
        )
        self.worker_thread.start()

        self.build_ui()
        self.poll_queue()

    def build_ui(self):
        self.root.title("RESISTOMAT 2316 - Messung")
        self.root.geometry("420x380")

        tk.Label(self.root, text="Aktueller Messwert:", font=("Segoe UI", 12)).pack(pady=(20, 0))
        self.value_label = tk.Label(self.root, text="--", font=("Segoe UI", 28, "bold"))
        self.value_label.pack(pady=10)

        self.progress_label = tk.Label(self.root, text=f"Messung 0 / {self.target_count}", font=("Segoe UI", 11))
        self.progress_label.pack(pady=(0, 10))

        self.confirm_button = tk.Button(
            self.root, text="Bestaetigen (Enter)", font=("Segoe UI", 12),
            command=self.confirm_measurement
        )
        self.confirm_button.pack(pady=10)

        tk.Label(self.root, text="Letzte Eintraege (Doppelklick zum Bearbeiten):", font=("Segoe UI", 10)).pack(pady=(10, 0))
        self.history_listbox = tk.Listbox(self.root, height=8, font=("Consolas", 10))
        self.history_listbox.pack(fill="both", expand=True, padx=20, pady=(0, 5))
        self.history_listbox.bind("<Double-Button-1>", lambda event: self.edit_selected_measurement())

        self.edit_button = tk.Button(
            self.root, text="Ausgewaehlten Wert bearbeiten", font=("Segoe UI", 10),
            command=self.edit_selected_measurement
        )
        self.edit_button.pack(pady=(0, 15))

        self.root.bind("<Return>", lambda event: self.confirm_measurement())

    def poll_queue(self):
        try:
            while True:
                raw_value = self.value_queue.get_nowait()
                self.current_raw_value = raw_value
                self.value_label.config(text=raw_value)
        except queue.Empty:
            pass

        if self.count < self.target_count:
            self.root.after(150, self.poll_queue)

    def confirm_measurement(self):
        if self.count >= self.target_count:
            return
        if self.current_raw_value is None:
            messagebox.showwarning("Kein Wert", "Es liegt noch kein Messwert vor.")
            return

        parsed = split_value_unit(self.current_raw_value)
        if parsed is None:
            messagebox.showwarning("Fehler", f"Konnte Wert nicht lesen: {self.current_raw_value}")
            return
        number, unit = parsed

        row = append_measurement(self.wb, self.ws, self.excel_path, number, unit)
        if row is None:
            messagebox.showerror(
                "Excel gesperrt",
                "Die Excel-Datei ist gerade geoeffnet (z.B. in Excel).\nBitte schliessen und erneut bestaetigen."
            )
            return

        self.count += 1
        self.history_rows.insert(0, row)
        self.history_raw.insert(0, self.current_raw_value)
        self.history_listbox.insert(0, f"#{self.count}: {self.current_raw_value}")
        self.progress_label.config(text=f"Messung {self.count} / {self.target_count}")

        if self.count >= self.target_count:
            self.confirm_button.config(state="disabled", text="Fertig!")
            self.value_label.config(text="Fertig")
            messagebox.showinfo("Fertig", f"Alle {self.target_count} Messungen erfasst.")
            self.stop_event.set()

    def edit_selected_measurement(self):
        selection = self.history_listbox.curselection()
        if not selection:
            messagebox.showinfo("Keine Auswahl", "Bitte zuerst einen Eintrag in der Liste auswaehlen.")
            return
        index = selection[0]

        new_raw = simpledialog.askstring(
            "Wert bearbeiten",
            "Neuer Messwert (z.B. 12.34 mOhm):",
            initialvalue=self.history_raw[index],
            parent=self.root,
        )
        if not new_raw:
            return

        parsed = split_value_unit(new_raw)
        if parsed is None:
            messagebox.showwarning("Fehler", f"Konnte Wert nicht lesen: {new_raw}")
            return
        number, unit = parsed

        row = self.history_rows[index]
        old_number, old_unit = self.ws.cell(row=row, column=2).value, self.ws.cell(row=row, column=3).value
        self.ws.cell(row=row, column=2, value=number)
        self.ws.cell(row=row, column=3, value=unit)
        try:
            self.wb.save(self.excel_path)
        except PermissionError:
            self.ws.cell(row=row, column=2, value=old_number)
            self.ws.cell(row=row, column=3, value=old_unit)
            messagebox.showerror(
                "Excel gesperrt",
                "Die Excel-Datei ist gerade geoeffnet (z.B. in Excel).\nBitte schliessen und erneut versuchen."
            )
            return

        self.history_raw[index] = new_raw
        display_count = self.count - index
        self.history_listbox.delete(index)
        self.history_listbox.insert(index, f"#{display_count}: {new_raw}")

    def on_close(self):
        self.stop_event.set()
        self.ser.close()
        self.root.destroy()


def main():
    root = tk.Tk()
    root.withdraw()

    target_count = simpledialog.askinteger(
        "Anzahl Messungen", "Wie viele Messungen sollen erfasst werden?",
        minvalue=1, initialvalue=30
    )
    if not target_count:
        return

    last_path = load_last_excel_path()
    if last_path:
        initialdir = str(Path(last_path).parent)
        initialfile = Path(last_path).name
    else:
        initialdir = str(Path.home())
        initialfile = "resistomat_messwerte.xlsx"

    excel_path = filedialog.asksaveasfilename(
        title="Excel-Datei waehlen/erstellen",
        defaultextension=".xlsx",
        filetypes=[("Excel-Datei", "*.xlsx")],
        initialdir=initialdir,
        initialfile=initialfile,
    )
    if not excel_path:
        return
    save_last_excel_path(excel_path)

    last_port = load_last_com_port()
    ser = None
    while ser is None:
        com_port = ask_com_port(root, last_port)
        if not com_port:
            return
        last_port = com_port

        try:
            ser = serial.Serial(
                port=com_port,
                baudrate=BAUDRATE,
                bytesize=serial.EIGHTBITS,
                parity=serial.PARITY_NONE,
                stopbits=serial.STOPBITS_ONE,
                timeout=5,
            )
        except serial.SerialException as e:
            messagebox.showerror("Verbindungsfehler", f"Konnte {com_port} nicht oeffnen:\n{e}")
            continue

        if not send(ser, "*idn?") or poll(ser) is None:
            messagebox.showerror("Geraetefehler", "Geraet antwortet nicht auf *idn? - Verbindung pruefen.")
            ser.close()
            ser = None

    save_last_com_port(com_port)

    root.deiconify()
    app = ResistomatUI(root, ser, excel_path, target_count)
    root.protocol("WM_DELETE_WINDOW", app.on_close)
    root.mainloop()


if __name__ == "__main__":
    main()
